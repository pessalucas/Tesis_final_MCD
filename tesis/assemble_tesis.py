# -*- coding: utf-8 -*-
"""Ensambla los capítulos markdown de tesis/ en un único .docx (borrador)."""
import re, glob, os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

BASE = os.path.dirname(os.path.abspath(__file__))
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
st.paragraph_format.space_after = Pt(6)

def add_runs(p, text):
    for i, part in enumerate(re.split(r"(\*\*.*?\*\*)", text)):
        if not part: continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        else:
            p.add_run(part)

def heading(text, level):
    h = doc.add_heading(level=level)
    for r in h.runs: r.font.color.rgb = RGBColor(0, 0, 0)
    add_runs(h, text)

def center(text, size, bold=False, after=6):
    p = doc.add_paragraph(); r = p.add_run(text); r.font.size = Pt(size); r.bold = bold
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(after)

# ---- carátula ----
for _ in range(3): doc.add_paragraph()
center("MAESTRÍA EN EXPLOTACIÓN DE DATOS Y GESTIÓN DEL CONOCIMIENTO", 14, True)
center("SEMINARIO DE TRABAJO FINAL — TESIS (BORRADOR)", 12, True, 24)
center("Modelos predictivos para la planificación de la producción en un cultivo hidropónico", 18, True, 24)
for _ in range(4): doc.add_paragraph()
center("Alumno: Lucas Pessagno", 12)
center("Director/a: [por confirmar]", 12)
center("Bariloche, 2026", 12)
doc.add_page_break()

# ---- abstract ----
heading("Resumen", 1)
RES = [
 ("¿Por qué lo hará?", "La planificación de la producción en cultivos hidropónicos depende hoy sobre todo de la experiencia del operador. Anticipar cuándo estará lista una cosecha, cuánto rendirá y qué condiciones influyen en el resultado es difícil, y la fuerte variabilidad estacional obliga a cosechas parciales que desaprovechan el potencial productivo."),
 ("¿Qué hará?", "El trabajo desarrolla, sobre datos reales de una operación hidropónica de albahaca, un conjunto de modelos predictivos con la estimación del tiempo de cosecha como eje central, re-estimable día a día durante el ciclo, más aristas exploratorias de proyección temprana y rendimiento."),
 ("¿Cómo lo hará?", "La estimación se encuadra como un problema de tiempo al evento mediante landmarking y modelos de aprendizaje automático, con interpretabilidad basada en SHAP y datos climáticos externos. La evaluación respeta la escala real de la muestra, validando por cosecha y de manera temporal, y reportando intervalos de cobertura calibrada."),
 ("¿Cuáles serán los beneficios?", "Los modelos permiten planificar entregas con mayor certeza y reducir las cosechas parciales. El trabajo aporta, además, un caso replicable de cómo obtener predicciones con rigor cuando los datos son valiosos pero limitados."),
]
for q, a in RES:
    p = doc.add_paragraph(); p.add_run(q + " ").bold = True; p.add_run(a)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p = doc.add_paragraph(); p.add_run("Palabras clave: ").bold = True
p.add_run("hidroponía; aprendizaje automático; tiempo de cosecha; análisis de tiempo al evento; agricultura de ambiente controlado; planificación de la producción.")
doc.add_page_break()

# ---- capítulos ----
def render_md(path):
    lines = open(path, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("> "):  # blockquote -> itálica
            p = doc.add_paragraph(); r = p.add_run(ln[2:]); r.italic = True
            i += 1; continue
        if ln.startswith("#"):
            lvl = len(ln) - len(ln.lstrip("#"))
            heading(ln[lvl:].strip(), min(lvl, 4)); i += 1; continue
        if ln.strip().startswith("|") and "|" in ln:
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i]); i += 1
            rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in block]
            rows = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]  # saca separador
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        if ci < len(t.rows[ri].cells):
                            cp = t.rows[ri].cells[ci].paragraphs[0]; add_runs(cp, cell)
                            if ri == 0:
                                for rr in cp.runs: rr.bold = True
            doc.add_paragraph(); continue
        if ln.strip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, ln.strip()[2:]); i += 1; continue
        if re.match(r"^\d+\.\s", ln.strip()):
            p = doc.add_paragraph(style="List Number"); add_runs(p, re.sub(r"^\d+\.\s", "", ln.strip())); i += 1; continue
        if ln.strip() == "":
            i += 1; continue
        p = doc.add_paragraph(); add_runs(p, ln); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; i += 1

for f in ["01_introduccion.md", "02_marco_teorico.md", "03_metodologia.md",
          "04_resultados.md", "05_discusion.md", "06_conclusiones.md", "07_referencias.md"]:
    render_md(os.path.join(BASE, f))
    doc.add_page_break()

out = os.path.join(BASE, "Tesis - borrador v1.docx")
doc.save(out)
print("Guardado:", out)
print("Párrafos:", len(doc.paragraphs))
